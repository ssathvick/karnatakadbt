from pathlib import Path
import math,json,textwrap
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parent; ASSETS=ROOT/'assets'; ASSETS.mkdir(exist_ok=True)
PROJECT=ROOT.parent
RED='#AD1C32';YELLOW='#F2C719';INK='#20202A';GREY='#6A6870';PALE='#FFF8DC'
def font(n,bold=False):return ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if bold else 'arial.ttf'),n)
def canvas(name,w=2100,h=850):
 im=Image.new('RGB',(w,h),'white');return im,ImageDraw.Draw(im),ASSETS/(name+'.png')
def txt(draw,xy,text,size=38,fill=INK,bold=False,width=32):
 lines=[]
 for line in text.split('\n'):lines+=textwrap.wrap(line,width=width) or ['']
 draw.multiline_text(xy,'\n'.join(lines),font=font(size,bold),fill=fill,spacing=10)
def arrow(d,a,b,color=RED,width=8):
 d.line([a,b],fill=color,width=width);theta=math.atan2(b[1]-a[1],b[0]-a[0]);p=[b,(b[0]-24*math.cos(theta-.5),b[1]-24*math.sin(theta-.5)),(b[0]-24*math.cos(theta+.5),b[1]-24*math.sin(theta+.5))];d.polygon(p,fill=color)
def flow(name,steps,notes=None):
 im,d,path=canvas(name,h=700)
 n=len(steps);gap=30;w=(2000-gap*(n-1))/n
 for i,(title,body) in enumerate(steps):
  x=50+i*(w+gap);d.ellipse((x+10,50,x+104,144),fill=RED if i%2==0 else YELLOW);txt(d,(x+42,70),str(i+1),44,'white' if i%2==0 else INK,True)
  txt(d,(x+10,185),title,44,bold=True,width=max(10,int(w/25)))
  txt(d,(x+10,310),body,40,width=max(12,int(w/22)))
  if i<n-1:arrow(d,(x+w-15,95),(x+w+20,95))
 if notes:txt(d,(60,590),notes,35,GREY,width=105)
 im.save(path);return path
flow('journey', [('Register','Create household and member records'),('Review','Validate evidence and return corrections'),('Connect','Synchronize approved records'),('Enroll','Apply programme rules and approve'),('Reconcile','Track simulated payments and exceptions')], 'Training environment: synthetic records and simulated payments. Registration approval is not a payment.')
flow('data-flow',[('Capture','Household data and location'),('Approve','Independent staff decision'),('Synchronize','Native registry and portal records'),('Confirm','Record identifiers and outcomes')], 'A saved record, an approved record and a synchronized record are different states.')
flow('payment-flow',[('Authorize','Approved programme and payment instruction'),('Validate','Financial address and account checks'),('Transfer','Submit an instruction once'),('Reconcile','Confirm final outcome or investigate')], 'Current partner journeys are simulations. Official state payment access requires separate onboarding.')
flow('knowledge-flow',[('Draft','Author writes task guidance'),('Review','Another authorised reviewer checks it'),('Publish','Supervisor publishes reviewed guidance'),('Learn','Search guidance and monitor usage')], 'AI can suggest text and explanations. Staff remain responsible for decisions and publication.')
flow('operations',[('Open controls','Start Portal Control Centre on the desktop'),('Start portal','Select Start for KarnatakaDBT'),('Wait for ready','Starting changes to Running after checks'),('Use or stop','Open the site or stop its application services')], 'Stopping the control centre does not stop the portal. Keep the hosting computer awake and connected.')
im,d,path=canvas('decision',h=870)
for x,y,label in [(80,100,'Programme rules'),(80,400,'Verified evidence'),(820,250,'Eligibility assessment'),(1560,100,'Eligible for review'),(1560,400,'Needs correction\nor not eligible')]:
 d.rounded_rectangle((x,y,x+450,y+180),radius=30,fill=PALE if x!=820 else '#FBE5E8',outline=YELLOW if x!=820 else RED,width=4);txt(d,(x+28,y+40),label,42,bold=True,width=18)
arrow(d,(550,190),(790,300));arrow(d,(550,490),(790,380));arrow(d,(1290,320),(1530,190));arrow(d,(1290,380),(1530,490));txt(d,(100,700),'Evidence unavailable or inconsistent -> review queue, not an automatic rejection.',34,width=100);im.save(path)
im,d,path=canvas('kutumba',h=1050)
steps=[('1  Family ID','Identify the family through the approved process'),('2  Permitted attributes','Fetch only fields needed by the programme'),('3  Compare and resolve','Match members and review discrepancies'),('4  Apply programme rules','Keep evidence separate from a decision')]
for i,(a,b) in enumerate(steps):
 col=i if i<2 else 3-i
 x=80+col*1030;y=65+(i//2)*430;d.rounded_rectangle((x,y,x+900,y+310),radius=22,fill=PALE,outline=YELLOW,width=4);txt(d,(x+35,y+35),a,43,bold=True,width=35);txt(d,(x+35,y+120),b,36,width=40)
arrow(d,(1000,220),(1080,220));arrow(d,(1540,390),(1540,470));arrow(d,(1080,650),(1000,650));txt(d,(85,930),'Proposed integration pattern. Current Kutumba access and API specifications must be approved.',31,GREY,width=110);im.save(path)
im,d,path=canvas('architecture',h=1130)
cols=[(60,'Official sources',['Kutumba','AHARA and FRUITS','Certificates and education','DBT payment platform']),(750,'Integration layer',['Approved source adapters','Identity and code crosswalks','Evidence with provenance','Retry and reconciliation queues']),(1450,'KarnatakaDBT',['Household registry','Review and eligibility','Programme enrolment','Analytics and audit'])]
for x,title,items in cols:
 txt(d,(x,50),title,43,bold=True,width=26)
 for j,item in enumerate(items):
  y=170+j*195;d.rounded_rectangle((x,y,x+570,y+140),radius=18,fill=PALE if x!=750 else '#FBE5E8');txt(d,(x+25,y+30),item,40,width=24)
for y in [235,430,625,820]:arrow(d,(650,y),(720,y));arrow(d,(1340,y),(1420,y))
txt(d,(65,1010),'Proposed architecture. Connectors stay disabled for live data until access is authorised.',32,GREY,width=110);im.save(path)
im,d,path=canvas('benefits',h=810)
items=[('Citizens','Fewer repeated entries\nClearer status and corrections'),('Staff','Focused review queues\nTraceable decisions'),('Programme owners','Consistent programme rules\nVisible exceptions'),('Operations teams','Independent portal controls\nPublic availability checks')]
for i,(a,b) in enumerate(items):
 x=70+(i%2)*1040;y=50+(i//2)*380;d.ellipse((x,y,x+75,y+75),fill=YELLOW);txt(d,(x+110,y+5),a,44,bold=True,width=30);txt(d,(x+110,y+100),b,36,width=39)
im.save(path)

doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
styles=doc.styles
for element in doc.element.iter(qn('w:pBdr')):element.getparent().remove(element)
for element in styles.element.iter(qn('w:pBdr')):element.getparent().remove(element)
for name,size in [('Normal',11),('Title',30),('Subtitle',14),('Heading 1',22),('Heading 2',14),('Heading 3',12)]:
 st=styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(8)
styles['Normal'].paragraph_format.line_spacing=1.1
styles['Subtitle'].font.italic=False
for name in ['Heading 1','Heading 2','Heading 3']:styles[name].paragraph_format.keep_with_next=True
styles['Caption'].font.size=Pt(9);styles['Caption'].font.color.rgb=RGBColor.from_string('555555')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('KarnatakaDBT  |  Training edition  |  ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='KarnatakaDBT platform training guide';doc.core_properties.author='Sathvick';doc.core_properties.subject='Software operation benefits and proposed external integrations'
def p(text='',boldlead=None):
 q=doc.add_paragraph()
 if boldlead and text.startswith(boldlead):q.add_run(boldlead).bold=True;q.add_run(text[len(boldlead):])
 else:q.add_run(text)
 return q
def h(text,level=1):doc.add_heading(text,level)
def page(title):doc.add_page_break();h(title)
def bullets(items):
 for item in items:doc.add_paragraph(item,style='List Bullet')
def steps(items):
 for i,item in enumerate(items,1):p(str(i)+'. '+item)
def pic(name,width=7,caption=None):
 path=Path(name) if Path(name).exists() else ASSETS/(name+'.png')
 q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(5)
 inline=q.add_run().add_picture(str(path),width=Inches(width));inline._inline.docPr.set('descr',caption or path.stem.replace('-',' '))
 if caption:doc.add_paragraph(caption,style='Caption')
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,head in enumerate(headers):t.rows[0].cells[i].text=head
 for row in rows:
  cells=t.add_row().cells
  for i,value in enumerate(row):cells[i].text=str(value)
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE' if ri==0 else ('FAFAFA' if ri%2==0 else 'FFFFFF'));pr.append(sh)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1.04
    for run in para.runs:run.font.size=Pt(10);run.bold=ri==0
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 p()
 return t
def link(title,url):
 q=doc.add_paragraph();hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),q.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'9E1B32');props.append(color);run.append(props);text=OxmlElement('w:t');text.text=title;run.append(text);hyp.append(run);q._p.append(hyp)

pic(PROJECT/'platform/assets/sathvick-black.png',1.05)
doc.add_paragraph('KarnatakaDBT platform\ntraining guide',style='Title')
doc.add_paragraph('Direct Benefit Transfer Fully AI Driven Automated Platform',style='Subtitle')
p('Built by Sathvick',boldlead='Built by Sathvick')
p('Training edition   16 October 2026   aidrivendbt.com')
pic('journey',7)
h('Purpose of this guide',2)
p('This guide explains how KarnatakaDBT organizes household records, staff review, programme enrolment, simulated payments and service insights. It is for trainers, registration staff, reviewers, programme supervisors and implementation teams.')
p('Learners should finish the guide able to follow a household through the software, explain what each status means, use the map and filters, and distinguish current demonstration features from proposed government connections.')
p('The current platform contains synthetic records and simulated financial journeys. Kutumba and other government integrations described here are proposed connections; no live access is implied. The platform does not issue real government benefits or payments.')

page('How to use the training guide')
p('Use the first half for an operator demonstration and the second half for integration and implementation discussions. The example programmes and amounts are training configurations, not official scheme rules.')
table(['Part','Pages','What learners should understand'],[
('Platform and people','3 to 4','Software modules, roles and household data entry'),('Operational journey','5 to 8','Review filters, synchronization, eligibility and simulated payments'),('Insights and guidance','9 to 10','District mapping, knowledge, AI assistance and analytics'),('External integration','11 to 14','Kutumba, other sources, access routes and connector design'),('Running and practising','15 to 17','Daily operation, exercises, benefits and production preparation'),('References','18 to 19','Official sources and terminology')],[1.65,.8,4.65])
h('Demonstration dataset',2)
table(['Measure','Training environment'],[('Households','1,240 fictional households'),('People','3,720 synthetic individual records'),('Geography','31 districts across four Karnataka revenue divisions'),('Distribution','40 households and 120 member records per district'),('Location markers','1,240 labelled simulated household points inside district polygons'),('Knowledge library','208 adapted published guides at the verified baseline')],[2,5.1])
p('A household record groups family members. A beneficiary is someone selected under a programme. Registration alone does not establish eligibility, entitlement, enrolment or payment.')
h('Suggested training sequence',2)
p('Allow about 90 minutes as a facilitator planning estimate: 15 minutes for orientation, 30 for the household journey, 15 for maps and guidance, 15 for external integrations and 15 for exercises. Adjust the timing to the audience.')

page('The software and its users')
p('KarnatakaDBT connects a citizen-facing experience with staff workflows and native backend services. Each module has a specific responsibility; a result in one module is not automatically a decision in another.')
table(['Capability','What it does'],[('Household and social registries','Store people, households, relationships and geographic attributes'),('Beneficiary and staff workspaces','Provide role-specific tasks and programme application views'),('Programme management and approval workflows','Define synthetic programmes and support recorded review and enrolment'),('Partner and payment simulation','Exercise identity, bank, treasury and settlement outcomes without real transfers'),('Maps and field delivery','Show district concentration, household coordinates and field-work demonstrations'),('Knowledge, AI and analytics','Explain tasks, support reviewed content and summarize operational activity'),('Shared services','Support identity, consent demonstrations, reference data, schemas and integration administration')],[2.45,4.65])
h('Roles in a training session',2)
table(['Role','Primary responsibility','Important boundary'],[('Citizen or assisted applicant','Maintain their own household and follow applications','Cannot browse or edit other households'),('Registration officer','Capture details and correct returned records','Does not independently approve their own work'),('Verification officer','Review evidence and approve or return a record','Follow assigned permissions and independent-review controls'),('Programme supervisor','Monitor queues, programme operations and exceptions','Registration review is separate from payment authorisation'),('Technical administrator','Operate services and support integrations','Technical access does not establish policy eligibility')],[1.5,2.9,2.7])
p('The demonstration administrator has broad permissions for walkthroughs. Production deployment should use individual accounts and responsibilities assigned to each role.')

page('Register a household and its location')
p('The citizen or assisted-registration journey captures a household, its members and its recorded district. The existing form supports up to eight members. Explain that all data entered during training must be fictional.')
steps(['Enter the household name, district, community or area, monthly income and member details.','Review names, dates of birth and other required fields before submission.','Enter optional latitude and longitude together, or use device location after browser permission.','Submit the registration and follow its review status in the workspace.'])
pic(PROJECT/'evidence/Beneficiary-latitude-longitude-entry.png',7,'Actual citizen workspace showing latitude and longitude entry for the household.')
p('Citizens can update their own household coordinates through the workspace, including after registration approval. Location changes do not change approval or benefit eligibility. Each household point is shared by its members.')
p('Coordinates must fall inside the recorded district. Orange map points are simulated; user-reported points are blue. Both remain unverified. A browser location is useful only when the device is at the intended household location.')

page('Review household registrations')
p('Open the staff page named Review connect enroll. The household table separates the review decision from the synchronization outcome. Use filters to work on a focused queue instead of scrolling through every record.')
pic(PROJECT/'evidence/Household-registration-filters.png',7,'Actual staff screen with Review and Synchronization dropdowns, reference search and a matching count.')
steps(['Select Approved, Rejected or Pending review in the Review dropdown.','Select Complete, Not started or another synchronization status present in the dataset.','Combine both selections and optionally search a household reference. Use Reset filters to return to all records.','Open Review details, check the household evidence and follow the permitted decision controls.'])
table(['Review status','Meaning','Next action'],[('Pending review','A registration awaits a decision','Review the latest evidence'),('Rejected','The case requires attention or correction','Read the reason and correct through the permitted workflow'),('Approved','Registration review is complete','Check synchronization and programme readiness')],[1.3,2.8,3])
p('Filter selections remain in the page URL and survive refresh. A zero-result message means no records match the combination; it is not a processing error. Filters change the displayed queue, not any household decision.')

page('Understand synchronization')
p('Synchronization connects an approved intake record with corresponding native registry and portal records. It helps services refer to the same household while retaining their own identifiers and responsibilities.')
pic('data-flow',7,'The record journey from captured information to confirmed native-system outcomes.')
table(['State or outcome','What staff should understand'],[('Not started','No completed synchronization is recorded yet; check the review stage and next permitted action'),('Pending or processing','Work is queued or running; avoid submitting duplicate requests'),('Complete','The recorded integration step succeeded; this does not mean a payment was made'),('Error or exception','Read the reason, correct the cause and use the supported retry or review path')],[1.8,5.3])
h('Location updates',2)
p('The geographic editor saves coordinates with source and verification information. Coordinate-only updates are audited and synchronized to the modern registry and both native Odoo registries. A private timer retries queued coordinate synchronization failures.')
h('What a trainer should emphasize',2)
bullets(['Keep review status and synchronization status separate. An approved registration may still need integration work.','Do not interpret a temporary service failure as proof that a household is ineligible.','Use the latest record version. The location editor rejects stale edits so an older form cannot silently overwrite a newer change.'])

page('Set up programmes and assess eligibility')
p('Programme management connects a reviewed household to a defined support programme. In the current demonstration, configurable rules include maximum monthly household income and minimum household size, alongside a simulated benefit amount.')
pic('decision',7,'Programme rules and evidence lead to an assessment that staff must interpret through the authorised workflow.')
h('Example only',2)
p('A trainer creates a fictional programme with a monthly-income ceiling of INR 20,000 and a minimum household size of three. A synthetic household with INR 8,500 monthly income and three members meets those two configured conditions. That assessment alone does not prove entitlement under an actual government scheme.')
steps(['Create or select the training programme and review its current configuration.','Submit a programme application for an appropriately reviewed household.','Run the supported eligibility review and inspect the returned result or exception.','Confirm enrolment only through the permitted programme action. Keep the decision traceable.'])
p('The demonstration catalogue includes Gruha Lakshmi, Anna Bhagya, social-security pension, farmer-support and student-scholarship examples. These are synthetic configurations and must not be presented as official eligibility rules or payment commitments.')

page('Follow simulated payments and exceptions')
p('The partner simulator demonstrates the financial journey and the problems a real integration must handle. It is separate from operational beneficiary and bank records and does not send money.')
pic('payment-flow',7,'A proposed production payment journey illustrated through the current simulation environment.')
table(['Scenario','Training lesson'],[('Successful transfer and settlement','Accepted instructions and final settlement are separate milestones'),('Identity mismatch or ineligible household','Resolve identity evidence or programme assessment before proceeding'),('Invalid account or insufficient funds','Validate the financial destination and authorised funding'),('Payment failure or delayed settlement','Track the provider outcome and avoid duplicate transfers'),('Reconciliation mismatch','Investigate a difference between instruction and settlement'),('Reversal','Record the reversal and reconcile the final financial position')],[2.55,4.55])
p('The simulator provides nine named scenarios: success, identity mismatch, ineligible household, invalid account, insufficient funds, payment failure, delayed settlement, reconciliation mismatch and reversal.')
p('For a live programme, the official state DBT and treasury or bank interfaces would need approved access, instruction identifiers, duplicate protection, settlement evidence and a defined reconciliation process. [8]')

page('Use the Karnataka district map')
p('Open District map or Geographic record data. Select individuals or households, then choose density or count. Density means registered records divided by approximate mapped district area; it is not population density or a confirmed-beneficiary rate.')
pic(PROJECT/'evidence/Karnataka-31-district-density-map.png',4.65,'Actual map showing 31 district boundaries, density colours and simulated household locations.')
p('Select a district to inspect its count and density. Click a household marker or search its reference to open coordinate entry. The seeded dataset has equal counts per district, so count mode uses the same colour while density varies with district area.')
p('Boundaries are simplified visualization data, not legal or cadastral boundaries. Bengaluru North and Bengaluru South aliases are normalized in the project catalogue; Vijayanagara is separate from Ballari. Map attribution is retained. [12, 16]')

page('Use guidance AI and analytics')
p('The knowledge library provides task guidance and contextual help. Learners can search by service or role, read workflow steps and distinguish locally implemented guidance from references describing other software releases.')
pic('knowledge-flow',7,'Guidance moves through authoring, independent review and publication before wider use.')
h('AI assistance',2)
p('The software can use a local model when available and a labelled retrieval fallback. AI can explain a workflow, help find guidance and suggest draft text. Suggestions require review; the model is not an authoritative source of household facts, scheme eligibility or payment approval.')
h('Operational analytics',2)
p('Analytics summarizes profiles, members, applications, review backlog, integration exceptions and guidance usage. Partner simulations have their own analytical context and should not be added to operational totals.')
h('Field and service demonstrations',2)
p('The wider training environment includes field-delivery, grievances, conditionality and partner scenarios. Use these to explain how an exception becomes a recorded task, how a reviewer resolves it, and how management sees the remaining backlog.')
bullets(['Check what a metric counts before interpreting it. Households and individuals are different units.','Read the source and timestamp of evidence. A fluent explanation does not make stale data current.','When guidance is wrong, correct and republish it through the review workflow rather than treating AI output as policy.'])

page('Connect Kutumba to the household journey')
p('Kutumba is the recommended first external integration for household verification. Official descriptions identify a family and individual repository with updates from source departments, entitlement-management capabilities and integrations with other government systems. [1]')
pic('kutumba',7,'Proposed Kutumba workflow for KarnatakaDBT. It is not a claim of current connectivity.')
p('The intended experience is to identify a family through an approved process, retrieve only permitted attributes, compare members and resolve discrepancies before applying programme rules. External evidence should keep its source, reference and retrieval time.')
p('Official 2022 documentation describes forward and reverse integration using REST APIs and XML or JSON. This demonstrates an integration pattern, not a current production contract. Obtain the latest specification, permitted fields, authentication and data-sharing conditions from the system owner. [2]')
p('Birth, death, address or household changes should become reviewable evidence updates. Missing or inconsistent information needs a correction route. Do not translate a family address into a verified household GPS coordinate or infer caste from a surname.')

page('Choose external verification sources')
p('The following uses are proposed for KarnatakaDBT. The sources themselves are documented government systems; the exact fields and interfaces accessible to this project remain subject to approval.')
table(['Source','Evidence and proposed use','Integration approach'],[
('Kutumba [1, 2]','Family and individual attributes; link household identity, prefill permitted details and review changes','Priority one. Departmental onboarding and current API specification'),
('AHARA and PDS [3, 4]','Ration-card and food-benefit records; validate attributes relevant to food-support programmes','Request the permitted departmental interface. Public reports are not an API agreement'),
('Nadakacheri and AJSK [5]','Income, caste and revenue certificates; verify references where a scheme requires them','Approved verification interface or an authorised document-verification workflow'),
('FRUITS [6]','Farmer registration, land and crop information and benefits received; support agricultural programme evidence','Scheme-specific departmental integration. NIC documents existing provider and user-application integrations'),
('Bhoomi [7]','Land records; support landholding checks when required by the programme','First confirm whether the necessary evidence is already available through FRUITS'),
('eJanMa [9]','Registered births and deaths; flag household lifecycle changes for review','Prefer an approved Kutumba feed when it provides sufficient evidence')],[1.35,3.15,2.6])
p('Use the minimum evidence needed for the programme. A source can be authoritative for one attribute while another department owns a different attribute. Conflicts should remain visible until resolved.')

page('Connect service payment and document systems')
table(['Source','Proposed role in KarnatakaDBT','Access position'],[
('Seva Sindhu [10]','Exchange application references and workflow status for a relevant departmental service','Agree the service and process with its department'),
('State Scholarship Portal and education systems [3, 11]','Support scholarship evidence and relevant existing-award checks','Separate education and scheme-specific access; do not assume all student data is available'),
('Official Karnataka DBT platform [8]','Identity or financial-address validation, authorised payment instructions and reconciliation','Dedicated onboarding workstream with the state platform and treasury or bank parties'),
('DigiLocker and API Setu [13, 14]','Retrieve supported certificates through an authorised requester flow','Documented partner application and approval route. It does not automatically grant Kutumba access'),
('Local Government Directory [12]','Map project districts, taluks, villages and wards to official location codes','Prepare an explicit crosswalk. Our project codes are not LGD codes'),
('Aadhaar authentication if required [15]','Provide an authorised authentication result for a permitted use case','Separate requesting-entity arrangement and current UIDAI requirements; portal login is insufficient')],[1.55,3,2.55])
p('Karnataka already has an official DBT platform between departmental beneficiary systems and treasury or banks. KarnatakaDBT should integrate with an authorised payment route rather than interpret its own demonstration payment screen as access to government funds. [8]')
p('Production access to Kutumba or state DBT payments is not established by public documentation or by operating aidrivendbt.com. The sponsoring department and system owners must confirm access, scope and current technical requirements.')

page('Design the external integration layer')
p('Recommended design: place provider-specific adapters between approved external sources and the software. Normalize identifiers and evidence before using them in a household review or programme decision.')
pic('architecture',7,'Proposed architecture separating official providers, controlled adapters and KarnatakaDBT workflows.')
table(['Evidence field','Purpose'],[('Provider and external reference','Identify the source system and the record being referenced'),('Attribute and permitted purpose','Explain which fact was retrieved and why'),('Retrieved at and source updated at','Distinguish retrieval time from the freshness of the underlying record'),('Verification outcome','Record verified, mismatch, unavailable, expired or needs-review outcomes'),('Internal record and mapping version','Link evidence to the correct household and geography codes'),('Decision reference','Record which authorised programme decision used that evidence')],[2.5,4.6])
p('Use separate demo, sandbox and production modes. Keep retry queues, duplicate-request protection and reconciliation outcomes explicit. External verification status must remain separate from the existing internal synchronization status.')

page('Start stop and operate the platform')
p('The portal is hosted through local application services and a dedicated Cloudflare tunnel. Public availability depends on the host computer, network, backend services and connector remaining available.')
pic('operations',7,'Daily operation through the local Portal Control Centre.')
steps(['After a Windows restart, double-click Start Portal Control Centre on the desktop.','Open the local controls at http://127.0.0.1:18212 and start KarnatakaDBT.','Wait while services warm up. The controls report Running after workload, connector and public HTTP checks pass.','Use Open to visit https://aidrivendbt.com. Use Stop to stop the selected portal when it is no longer needed.'])
h('Read the status before troubleshooting',2)
table(['Status or symptom','Operator response'],[('Starting','Allow the backend to warm up; avoid repeated start requests'),('Running','Public checks and application readiness passed at the last check'),('Checking or Degraded','Wait for a fresh check and inspect the reported issue'),('Error or Cloudflare 1033','Check the runtime and tunnel; a saved old status does not prove availability'),('Control page closed','The portal continues running; reopen the controls when needed')],[2,5.1])
p('Start Portal Control Centre and Stop Portal Control Centre manage the controls service. They are separate from the portal Start and Stop buttons. Keep the hosting computer awake and connected when the public site must remain available.')

page('Run a practical training exercise')
p('Use a fictional household and a training programme. Do not enter real government identifiers, banking details or personal records into the demonstration.')
steps(['Citizen view: open the household workspace and identify the household reference, members, district and current status.','Location: inspect latitude and longitude. Explain whether the point is simulated or user-reported. Enter a valid pair for a disposable training household.','Staff queue: filter Review to Pending review and Synchronization to Not started. Explain a zero-result combination if one occurs.','Independent review: inspect a permitted training record and its evidence. Discuss whether it should be returned or approved; only take a decision authorised for the exercise.','Synchronization: locate an approved record and explain the difference between Approved and Complete.','Programme journey: review the configured income and household-size rules, run a permitted assessment and inspect the result.','Simulation: run a delayed-settlement or invalid-account scenario. Identify the next exception-handling step.','Map and guidance: select a district, compare count with density and find the relevant knowledge article.'])
h('Trainer questions and expected answers',2)
table(['Question','Expected answer'],[('Does Approved mean paid','No. It is a registration-review status'),('Does Complete mean eligible','No. It describes the recorded synchronization step'),('Why do equal counts have different density colours','District areas differ'),('Can a citizen edit another household location','No. Their write access is scoped to their own household'),('Can we call Kutumba simply because the site is online','No. Approved access and a current integration contract are still needed'),('What if external evidence is unavailable','Record the exception and use the defined review or correction path')],[3.1,4])

page('Explain benefits and prepare for production')
pic('benefits',7,'Expected operational benefits when workflows, data quality and responsibilities are maintained.')
p('The practical benefits are reduced repeated entry, clearer review queues, consistent evidence handling, traceable decisions and better visibility of exceptions. Treat these as expected benefits to measure, not guaranteed savings or proven reductions in fraud.')
h('Recommended first integration pilot',2)
p('Start with Kutumba household verification, certificate evidence and one programme. Keep payments simulated until the approved financial integration is ready. Demonstrate how a mismatch, missing record and source outage are handled before expanding the pilot.')
table(['Preparation area','Completion evidence'],[('Government access','Sponsoring department, permitted purpose, approved fields and current API contracts'),('Identity and access','Individual staff accounts, appropriate authentication and role permissions; replace shared demonstration passwords'),('Data environment','Separate production from synthetic records, credentials and simulated payments'),('Operational readiness','Stable hosting, backups, monitoring, support ownership and recovery exercises'),('Programme governance','Versioned rules, authorised decisions, correction routes and payment reconciliation responsibilities')],[2,5.1])
p('Pilot checks: correct matching, source provenance, safe retries, duplicate prevention, usable corrections and explainable decisions. Distinguish AI assistance from authoritative evidence.')

sources=[
('1','Kutumba overview and integration context','National e-Governance Division','https://negd.gov.in/isl/Directory/statedata/403'),
('2','Kutumba implementation and API integration description','DARPG publication July to December 2022','https://darpg.gov.in/sites/default/files/MGMG-Jul-Dec%202022.pdf'),
('3','Karnataka state systems including AHARA and SSP','National Informatics Centre','https://www.nic.gov.in/state-content/karnataka/'),
('4','Karnataka public distribution system links','National Food Security Portal','https://nfsa.gov.in/State/KA'),
('5','Nadakacheri AJSK services and certificate verification','Bagalkote district administration','https://bagalkot.nic.in/en/service/nadakacheri-ajsk/'),
('6','FRUITS capabilities and integrations','National Informatics Centre','https://www.nic.gov.in/project/fruits/'),
('7','Bhoomi land records system','Bengaluru Urban district administration','https://bengaluruurban.nic.in/en/bhoomi/'),
('8','Karnataka DBT platform and state projects','NIC Informatics April 2024','https://informatics.nic.in/uploads/pdfs/6559c107_informatics_apr_2024.pdf'),
('9','eJanMa birth and death registration integrations','NIC Informatics January 2025','https://informatics.nic.in/files/websites/january-2025/karnatakas-eJanma-revolutionizes-issuance-of-digitally-signed-birth-and-death-certificates.php'),
('10','Seva Sindhu integration and citizen service context','DARPG regional conference proceedings','https://darpg.gov.in/sites/default/files/Srinagar%20RC%20Proceedings.pdf'),
('11','State Scholarship Portal overview','NIC Karnataka state projects','https://informatics.nic.in/files/websites/july-2024/karnataka.php'),
('12','Local Government Directory and standard location codes','Ministry of Panchayati Raj','https://lgdirectory.gov.in/'),
('13','DigiLocker partner documentation and specifications','API Setu','https://apisetu.gov.in/digilocker'),
('14','API consumer subscription and publisher approval','API Setu documentation','https://docs.apisetu.gov.in/document-central/explore-apisetu/Overview.html'),
('15','Authentication requesting entities and onboarding','Unique Identification Authority of India','https://uidai.gov.in/hi/authentication-requesting-agency'),
('16','District visualization data and OpenStreetMap attribution','Vardhan Systems and OpenStreetMap contributors','https://github.com/Vardhan-Systems/vardhan-maps')]
page('Official references for verification')
p('External-source descriptions were reviewed on 16 October 2026. Older publications establish historical integration patterns; confirm current endpoints, fields, access conditions and specifications with the provider before implementation. Reference numbers in the guide point to the sources below.')
for number,title,owner,url in sources[:8]:
 link(number+'  '+title,url);p(owner)
p('This guide describes the current synthetic KarnatakaDBT implementation and a proposed integration approach. It is not an official Government of Karnataka training publication or a statement of government endorsement.')
page('Further references and terminology')
for number,title,owner,url in sources[8:]:link(number+'  '+title,url);p(owner)
h('Key terms',2)
p('Household: a group of recorded members. Registry: the system holding structured records. Eligibility: an assessment against programme rules. Enrolment: admission to a programme through its permitted workflow. Synchronization: exchange or alignment of records between systems. Reconciliation: comparison of instructions and final outcomes. Provenance: where evidence came from and when it was obtained. Kutumba Family ID: an external family identifier, distinct from this project\'s synthetic household reference.')
p('KarnatakaDBT  |  Built by Sathvick  |  https://aidrivendbt.com')

out=ROOT/'KarnatakaDBT_Training_Guide.docx';doc.save(out)
print(str(out));print('Planned pages: 19; images:',len(doc.inline_shapes))
